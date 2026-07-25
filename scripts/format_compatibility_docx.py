"""Apply a compact landscape layout to the DP-800 compatibility matrix."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


PATH = Path("output/doc/DP-800-Matriz-Compatibilidade-Plataformas.docx")


def set_cell_text_size(cell, size: float) -> None:
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(size)


document = Document(PATH)
section = document.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(11.69)
section.page_height = Inches(8.27)
section.top_margin = Inches(0.45)
section.bottom_margin = Inches(0.45)
section.left_margin = Inches(0.4)
section.right_margin = Inches(0.4)

styles = document.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9)
for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
    styles[style_name].font.name = "Aptos Display"
    styles[style_name].font.size = Pt(size)

for paragraph in document.paragraphs:
    if paragraph.style.name == "Title":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

for table in document.tables:
    table.style = "Table"
    table.autofit = True
    for row_number, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_text_size(cell, 7.3 if len(table.columns) >= 7 else 8)
            if row_number == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True

document.save(PATH)
print(PATH)
